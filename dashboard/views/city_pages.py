"""
City Pages dashboard views — single create, bulk CSV/Excel import, manage templates.
"""
import csv
import io
import json
from django.conf import settings
from django.db import transaction
from django.shortcuts import render, get_object_or_404, redirect
from django.views import View
from django.contrib import messages
from django.http import HttpResponse
from django.utils.text import slugify

import openpyxl

from accounts.permissions import SEOManagerRequiredMixin
from dashboard.mixins import LoggedActionMixin
from pages.models import PageTemplate, CityPage
from catalog.models import Product
from media_library.models import MediaAsset


def _resolve_slug(raw_slug, template, city, state, *, exclude_pk=None):
    """A staff-entered slug wins (slugified); otherwise fall back to the
    template's auto pattern, exactly like before this field existed.
    Returns (slug, error) — error is a user-facing string, or None."""
    custom = slugify((raw_slug or "").strip())
    slug = custom or template.build_slug(city, state)
    qs = CityPage.objects.filter(slug=slug)
    if exclude_pk:
        qs = qs.exclude(pk=exclude_pk)
    if qs.exists():
        return slug, f"URL slug “{slug}” is already used by another city page — choose a different one."
    return slug, None


def _resolve_canonical(raw_canonical, slug):
    """A staff-entered canonical URL wins verbatim; otherwise auto-build it
    from the production site origin + this page's slug."""
    custom = (raw_canonical or "").strip()
    if custom:
        return custom
    return f"{settings.SITE_URL.rstrip('/')}/{slug}/"


class CityPageIndexView(SEOManagerRequiredMixin, View):
    def get(self, request):
        return render(request, "dashboard/city_pages/index.html", {
            "active_nav": "city_pages",
        })


class CityPageListView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request):
        qs = CityPage.objects.select_related("template").order_by("city")
        template_id = request.GET.get("template", "")
        state       = request.GET.get("state", "")
        status      = request.GET.get("status", "")
        q           = request.GET.get("q", "")
        if template_id:
            qs = qs.filter(template_id=template_id)
        if state:
            qs = qs.filter(state__icontains=state)
        if status:
            qs = qs.filter(status=status)
        if q:
            qs = qs.filter(city__icontains=q) | qs.filter(slug__icontains=q)
        return render(request, "dashboard/city_pages/list.html", {
            "city_pages": qs,
            "templates":  PageTemplate.objects.all(),
            "template_id": template_id, "state": state, "status": status, "q": q,
            "active_nav": "city_pages",
        })

    def post(self, request):
        action = request.POST.get("bulk_action")
        ids    = request.POST.getlist("page_ids")
        if not ids or action not in ("publish", "draft", "delete"):
            messages.error(request, "Select at least one page and a valid action.")
            return redirect("city_page_list")
        qs    = CityPage.objects.filter(pk__in=ids)
        count = qs.count()
        if action == "publish":
            qs.update(status="published")
            self.log_action(f"Bulk published {count} city page(s)")
            messages.success(request, f"{count} page(s) published.")
        elif action == "draft":
            qs.update(status="draft")
            self.log_action(f"Bulk set {count} city page(s) to draft")
            messages.success(request, f"{count} page(s) set to draft.")
        elif action == "delete":
            qs.delete()
            self.log_action(f"Bulk deleted {count} city page(s)")
            messages.success(request, f"{count} page(s) deleted.")
        return redirect("city_page_list")


class CityPageCreateView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request):
        tpl_id = request.GET.get("template")
        template = PageTemplate.objects.filter(pk=tpl_id).first() if tpl_id else None
        return render(request, "dashboard/city_pages/form.html", {
            "templates": PageTemplate.objects.all(),
            "products":  Product.objects.filter(status="published"),
            "selected_template": template,
            "schema_choices": CityPage._meta.get_field("schema_type").choices,
            "site_url": settings.SITE_URL.rstrip("/"),
            "active_nav": "city_pages",
        })

    def post(self, request):
        d = request.POST
        template = get_object_or_404(PageTemplate, pk=d["template"])
        city, state = d["city"], d.get("state", "")
        slug, slug_error = _resolve_slug(d.get("slug"), template, city, state)
        if slug_error:
            messages.error(request, slug_error)
            return redirect(f"{request.path}?template={template.pk}")

        page = CityPage(
            template        = template,
            city            = city,
            state           = state,
            slug            = slug,
            h1_title        = d.get("h1_title") or None,
            hero_heading    = d.get("hero_heading") or None,
            hero_description = d.get("hero_description") or None,
            main_content    = d.get("main_content") or None,
            why_choose_us   = json.loads(d.get("why_choose_us_json", "null")),
            faqs            = json.loads(d.get("faqs_json", "null")),
            testimonials    = json.loads(d.get("testimonials_json", "[]")),
            seo_title       = d.get("seo_title", ""),
            meta_description = d.get("meta_description", ""),
            meta_keywords   = d.get("meta_keywords", ""),
            canonical_url   = _resolve_canonical(d.get("canonical_url"), slug),
            og_title        = d.get("og_title", ""),
            og_description  = d.get("og_description", ""),
            twitter_title   = d.get("twitter_title", ""),
            twitter_description = d.get("twitter_description", ""),
            schema_type     = d.get("schema_type", ""),
            status          = d.get("status", "draft"),
        )
        page.save()
        rel_ids = d.getlist("related_products")
        if rel_ids:
            page.related_products.set(rel_ids)
        self.log_action("Created city page", page)
        messages.success(request, f"City page '{page}' created.")
        return redirect("city_page_list")


class CityPageEditView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request, pk):
        page = get_object_or_404(CityPage, pk=pk)
        resolved = page.resolved()
        return render(request, "dashboard/city_pages/form.html", {
            "page":      page,
            "resolved":  resolved,
            "templates": PageTemplate.objects.all(),
            "products":  Product.objects.filter(status="published"),
            "selected_template": page.template,
            "schema_choices": CityPage._meta.get_field("schema_type").choices,
            "site_url": settings.SITE_URL.rstrip("/"),
            "active_nav": "city_pages",
        })

    def post(self, request, pk):
        page = get_object_or_404(CityPage, pk=pk)
        d = request.POST
        template = get_object_or_404(PageTemplate, pk=d["template"])
        city, state = d["city"], d.get("state", "")
        slug, slug_error = _resolve_slug(d.get("slug"), template, city, state, exclude_pk=page.pk)
        if slug_error:
            messages.error(request, slug_error)
            return redirect("city_page_edit", pk=pk)

        page.template        = template
        page.city            = city
        page.state           = state
        page.slug             = slug
        page.h1_title        = d.get("h1_title") or None
        page.hero_heading    = d.get("hero_heading") or None
        page.hero_description = d.get("hero_description") or None
        page.main_content    = d.get("main_content") or None
        page.why_choose_us   = json.loads(d.get("why_choose_us_json", "null"))
        page.faqs            = json.loads(d.get("faqs_json", "null"))
        page.testimonials    = json.loads(d.get("testimonials_json", "[]"))
        page.seo_title       = d.get("seo_title", "")
        page.meta_description = d.get("meta_description", "")
        page.meta_keywords   = d.get("meta_keywords", "")
        page.canonical_url   = _resolve_canonical(d.get("canonical_url"), slug)
        page.og_title        = d.get("og_title", "")
        page.og_description  = d.get("og_description", "")
        page.twitter_title   = d.get("twitter_title", "")
        page.twitter_description = d.get("twitter_description", "")
        page.schema_type     = d.get("schema_type", "")
        page.status          = d.get("status", "draft")
        page.save()
        rel_ids = d.getlist("related_products")
        page.related_products.set(rel_ids)
        self.log_action("Updated city page", page)
        messages.success(request, "City page updated.")
        return redirect("city_page_list")


class CityPageDeleteView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def post(self, request, pk):
        page = get_object_or_404(CityPage, pk=pk)
        self.log_action("Deleted city page", page)
        page.delete()
        messages.success(request, "City page deleted.")
        return redirect("city_page_list")


class CityPagePreviewView(SEOManagerRequiredMixin, View):
    """Sends the editor to the real Next.js city-page route instead of a
    hand-built mock — see ProductPreviewView (dashboard/views/products.py)
    for the same pattern and why."""
    def get(self, request, pk):
        from django.conf import settings
        from api.preview_tokens import make_preview_token

        page = get_object_or_404(CityPage, pk=pk)
        token = make_preview_token("citypage", page.slug)
        return redirect(f"{settings.FRONTEND_URL}/{page.slug}?preview={token}")


# ── Bulk Import ───────────────────────────────────────────────────────────────

REQUIRED_COLS = ["city", "state", "product_type"]

def _parse_rows(file_obj, filename):
    """Parse CSV or XLSX; return list of dicts."""
    rows = []
    fname = filename.lower()
    if fname.endswith(".xlsx"):
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(headers, row)))
    else:
        text    = file_obj.read().decode("utf-8-sig")
        reader  = csv.DictReader(io.StringIO(text))
        rows    = list(reader)
    return rows


class BulkImportView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    template_name = "dashboard/city_pages/bulk_import.html"

    def get(self, request):
        return render(request, self.template_name, {
            "templates": PageTemplate.objects.all(),
            "active_nav": "city_pages",
        })

    def post(self, request):
        action = request.POST.get("action", "preview")

        if action == "download_template":
            return self._download_template()

        if action == "commit":
            raw_rows = json.loads(request.POST.get("raw_rows", "[]"))
            created  = 0
            with transaction.atomic():
                for row in raw_rows:
                    city    = str(row.get("city") or "").strip()
                    state   = str(row.get("state") or "").strip()
                    product = str(row.get("product_type") or "").strip()
                    tpl = PageTemplate.objects.filter(product_type_label__iexact=product).first()
                    if not tpl or not city:
                        continue
                    slug = tpl.build_slug(city, state)
                    if CityPage.objects.filter(slug=slug).exists():
                        continue
                    page = CityPage(
                        template     = tpl,
                        city         = city,
                        state        = state,
                        h1_title     = row.get("h1_title") or None,
                        seo_title    = row.get("seo_title", ""),
                        meta_description = row.get("meta_description", ""),
                        status       = "draft",
                    )
                    page.save()
                    created += 1
            self.log_action(f"Bulk imported {created} city pages")
            messages.success(request, f"{created} city pages created.")
            return redirect("city_page_list")

        # action == "preview" — requires an uploaded file
        uploaded = request.FILES.get("import_file")
        if not uploaded:
            messages.error(request, "No file uploaded.")
            return redirect("city_page_bulk_import")

        rows  = _parse_rows(uploaded, uploaded.name)
        errors, preview_rows = [], []

        for i, row in enumerate(rows, start=2):
            row_errors = []
            city    = str(row.get("city") or "").strip()
            state   = str(row.get("state") or "").strip()
            product = str(row.get("product_type") or "").strip()
            if not city:
                row_errors.append("Missing city")
            if not product:
                row_errors.append("Missing product_type")
            tpl = PageTemplate.objects.filter(product_type_label__iexact=product).first()
            if not tpl:
                row_errors.append(f"No template matching '{product}'")
            slug = tpl.build_slug(city, state) if tpl else ""
            dupe = CityPage.objects.filter(slug=slug).exists() if slug else False
            if dupe:
                row_errors.append(f"Slug '{slug}' already exists")
            preview_rows.append({
                "row": i, "city": city, "state": state, "product": product,
                "slug": slug, "errors": row_errors,
            })
            if row_errors:
                errors.append((i, row_errors))

        return render(request, self.template_name, {
            "preview_rows": preview_rows,
            "has_errors":   bool(errors),
            "raw_rows":     json.dumps(rows),
            "templates":    PageTemplate.objects.all(),
            "active_nav":   "city_pages",
        })

    def _download_template(self):
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="city_pages_template.csv"'
        writer = csv.writer(resp)
        writer.writerow(["city", "state", "product_type", "h1_title", "hero_heading", "meta_description"])
        writer.writerow(["Ahmedabad", "Gujarat", "Laminates", "", "", ""])
        return resp


# ── Page Templates ────────────────────────────────────────────────────────────

class PageTemplateListView(SEOManagerRequiredMixin, View):
    def get(self, request):
        return render(request, "dashboard/city_pages/template_list.html", {
            "templates":  PageTemplate.objects.all(),
            "active_nav": "city_pages",
        })


class PageTemplateCreateView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request):
        return render(request, "dashboard/city_pages/template_form.html", {
            "products":   Product.objects.filter(status="published"),
            "schema_choices": CityPage._meta.get_field("schema_type").choices,
            "active_nav": "city_pages",
        })

    def post(self, request):
        d = request.POST
        tpl = PageTemplate.objects.create(
            name=d["name"],
            product_type_label=d["product_type_label"],
            url_pattern=d.get("url_pattern", "{product}-in-{city}"),
            default_h1=d.get("default_h1", ""),
            default_hero_heading=d.get("default_hero_heading", ""),
            default_hero_desc=d.get("default_hero_desc", ""),
            default_main_content=d.get("default_main_content", ""),
            why_choose_us=json.loads(d.get("why_choose_us_json", "[]")),
            faqs=json.loads(d.get("faqs_json", "[]")),
            default_schema_type=d.get("default_schema_type", "LocalBusiness"),
        )
        rel_ids = d.getlist("related_products")
        if rel_ids:
            tpl.related_products.set(rel_ids)
        self.log_action("Created page template", tpl)
        messages.success(request, f"Template '{tpl.name}' created.")
        return redirect("page_template_list")


class PageTemplateEditView(SEOManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request, pk):
        tpl = get_object_or_404(PageTemplate, pk=pk)
        return render(request, "dashboard/city_pages/template_form.html", {
            "template":  tpl,
            "products":  Product.objects.filter(status="published"),
            "schema_choices": CityPage._meta.get_field("schema_type").choices,
            "active_nav": "city_pages",
        })

    def post(self, request, pk):
        tpl = get_object_or_404(PageTemplate, pk=pk)
        d = request.POST
        tpl.name = d["name"]
        tpl.product_type_label = d["product_type_label"]
        tpl.url_pattern = d.get("url_pattern", tpl.url_pattern)
        tpl.default_h1 = d.get("default_h1", "")
        tpl.default_hero_heading = d.get("default_hero_heading", "")
        tpl.default_hero_desc = d.get("default_hero_desc", "")
        tpl.default_main_content = d.get("default_main_content", "")
        tpl.why_choose_us = json.loads(d.get("why_choose_us_json", "[]"))
        tpl.faqs = json.loads(d.get("faqs_json", "[]"))
        tpl.default_schema_type = d.get("default_schema_type", tpl.default_schema_type)
        tpl.save()
        rel_ids = d.getlist("related_products")
        tpl.related_products.set(rel_ids)
        self.log_action("Updated page template", tpl)
        messages.success(request, "Template updated.")
        return redirect("page_template_list")
