import csv
import io
import json
import os
import re
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.shortcuts import render, get_object_or_404, redirect
from django.views import View
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.utils.text import slugify

PRODUCTS_PER_PAGE = 20

import openpyxl

from accounts.permissions import ContentManagerRequiredMixin
from dashboard.mixins import LoggedActionMixin
from catalog.models import (
    Product, Category, Collection, ProductImage, ProductAttributeOption,
)
from media_library.models import MediaAsset, MediaFolder

MANAGED_ATTRIBUTES = {c[0] for c in ProductAttributeOption.ATTRIBUTE_CHOICES}


def _attribute_options(attribute):
    """[{id, value}, ...] in display order, for a managed dropdown field."""
    return list(
        ProductAttributeOption.objects.filter(attribute=attribute)
        .order_by("position", "id")
        .values("id", "value")
    )


def _attr_context():
    """Context vars consumed by the product form template's dropdowns."""
    return {
        f"{attribute}_options": _attribute_options(attribute)
        for attribute in MANAGED_ATTRIBUTES
    }


def _attr_lookup(attribute):
    """{lowercased value: canonical value} for case-insensitive CSV import matching."""
    return {
        o["value"].lower(): o["value"]
        for o in ProductAttributeOption.objects.filter(attribute=attribute).values("value")
    }

# Fields whose display on the product page can be toggled independently of
# their value — a checkbox posts only when checked, so absence means "off".
_VISIBILITY_FIELDS = [
    "show_surface", "show_product_type", "show_finish", "show_surface_category",
    "show_thickness", "show_dimensions", "show_application", "show_design_type",
]


def _resolve_status(data, has_images, request):
    """Read the requested status off the form, but refuse to publish a product
    that has no gallery images — it would show the placeholder image on the
    storefront instead of a real photo."""
    status = data.get("status", "draft")
    if status == Product.STATUS_PUBLISHED and not has_images:
        messages.error(
            request,
            "Can't publish without at least one image — the product would show the "
            "placeholder image. Add an image or save as Draft.",
        )
        return Product.STATUS_DRAFT
    return status


def _save_product_images(product, data):
    """Rebuild `product.product_images` from the form's three image panels:
    the gallery/swatch list, the single application (room) image, and the
    texture-variant repeater."""
    product.product_images.all().delete()
    next_pos = 0

    image_ids = [i for i in data.getlist("image_ids") if i.strip()]
    for asset_id in image_ids:
        try:
            asset = MediaAsset.objects.get(pk=int(asset_id))
        except (MediaAsset.DoesNotExist, ValueError, TypeError):
            continue
        ProductImage.objects.create(
            product=product, asset=asset, position=next_pos, role=ProductImage.ROLE_GALLERY,
        )
        next_pos += 1

    app_id = (data.get("application_image_id") or "").strip()
    if app_id:
        try:
            asset = MediaAsset.objects.get(pk=int(app_id))
        except (MediaAsset.DoesNotExist, ValueError, TypeError):
            asset = None
        if asset:
            ProductImage.objects.create(
                product=product, asset=asset, position=next_pos, role=ProductImage.ROLE_APPLICATION,
            )
            next_pos += 1

    try:
        texture_items = json.loads(data.get("texture_images_json", "[]"))
    except json.JSONDecodeError:
        texture_items = []
    for item in texture_items:
        try:
            asset = MediaAsset.objects.get(pk=int(item.get("id")))
        except (MediaAsset.DoesNotExist, ValueError, TypeError):
            continue
        ProductImage.objects.create(
            product=product, asset=asset, position=next_pos,
            role=ProductImage.ROLE_TEXTURE, label=(item.get("label") or "").strip(),
        )
        next_pos += 1


class ProductListView(ContentManagerRequiredMixin, View):
    def get(self, request):
        qs = Product.objects.select_related("category", "collection").order_by("-created")
        q  = request.GET.get("q", "")
        cat = request.GET.get("category", "")
        collection = request.GET.get("collection", "")
        status = request.GET.get("status", "")
        if q:
            qs = qs.filter(name__icontains=q) | qs.filter(sku__icontains=q)
        if cat:
            qs = qs.filter(category_id=cat)
        if collection:
            qs = qs.filter(collection_id=collection)
        if status:
            qs = qs.filter(status=status)

        total_count = qs.count()
        paginator = Paginator(qs, PRODUCTS_PER_PAGE)
        page_obj = paginator.get_page(request.GET.get("page"))
        elided_pages = list(paginator.get_elided_page_range(
            page_obj.number, on_each_side=2, on_ends=1
        ))

        querystring = request.GET.copy()
        querystring.pop("page", None)

        return render(request, "dashboard/products/list.html", {
            "products":    page_obj,
            "page_obj":    page_obj,
            "paginator":   paginator,
            "elided_pages": elided_pages,
            "total_count": total_count,
            "querystring": querystring.urlencode(),
            "categories":  Category.objects.all(),
            "collections": Collection.objects.all(),
            "q": q, "selected_cat": cat, "selected_collection": collection, "selected_status": status,
            "active_nav": "products",
        })


class ProductCreateView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request):
        return render(request, "dashboard/products/form.html", {
            "categories":  Category.objects.filter(status="published"),
            "collections": Collection.objects.filter(status="published"),
            "all_products": Product.objects.all(),
            "active_nav": "products",
            **_attr_context(),
        })

    def post(self, request):
        data = request.POST
        image_ids = [i for i in data.getlist("image_ids") if i.strip()]
        image_urls = json.loads(data.get("image_urls_json", "[]"))
        has_images = bool(image_ids) or bool(image_urls)
        product = Product(
            name=data["name"],
            sku=data["sku"],
            category_id=data["category"],
            collection_id=data.get("collection") or None,
            short_description=data.get("short_description", ""),
            description=data.get("description", ""),
            features=json.loads(data.get("features_json", "[]")),
            tech_specs=json.loads(data.get("specs_json", "{}")),
            image_urls=image_urls,
            finish=data.get("finish", ""),
            thickness=data.get("thickness", ""),
            dimensions=data.get("dimensions", ""),
            surface=data.get("surface", ""),
            product_type=data.get("product_type", ""),
            surface_category=data.get("surface_category", ""),
            application=data.get("application", ""),
            design_type=data.get("design_type", ""),
            color=data.get("color", ""),
            badge=data.get("badge", ""),
            accent_color=data.get("accent_color", "") or "#85addc",
            meta_title=data.get("meta_title", ""),
            meta_description=data.get("meta_description", ""),
            meta_keywords=data.get("meta_keywords", ""),
            status=_resolve_status(data, has_images, request),
            **{f: (f in data) for f in _VISIBILITY_FIELDS},
        )
        product.application_image_url = data.get("application_image_url", "")
        product.save()
        _save_product_images(product, data)
        # Related products — filter out blank strings
        rel_ids = [r for r in request.POST.getlist("related_products") if r.strip()]
        if rel_ids:
            product.related_products.set(rel_ids)
        self.log_action("Created product", product)
        messages.success(request, f"Product '{product.name}' created.")
        return redirect("product_edit", pk=product.pk)


class ProductEditView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    def get(self, request, pk):
        product = get_object_or_404(Product, pk=pk)
        return render(request, "dashboard/products/form.html", {
            "product":     product,
            "categories":  Category.objects.filter(status="published"),
            "collections": Collection.objects.filter(status="published"),
            "all_products": Product.objects.exclude(pk=pk),
            "current_images": product.product_images.filter(role=ProductImage.ROLE_GALLERY).select_related("asset"),
            "current_application_image": product.product_images.filter(
                role=ProductImage.ROLE_APPLICATION
            ).select_related("asset").first(),
            "current_texture_images": product.product_images.filter(
                role=ProductImage.ROLE_TEXTURE
            ).select_related("asset"),
            "active_nav": "products",
            **_attr_context(),
        })

    def post(self, request, pk):
        product = get_object_or_404(Product, pk=pk)
        data = request.POST
        product.name              = data["name"]
        product.sku               = data["sku"]
        product.category_id       = data["category"]
        product.collection_id     = data.get("collection") or None
        product.short_description = data.get("short_description", "")
        product.description       = data.get("description", "")
        product.features       = json.loads(data.get("features_json", "[]"))
        product.tech_specs     = json.loads(data.get("specs_json", "{}"))
        image_urls = json.loads(data.get("image_urls_json", "[]"))
        product.image_urls     = image_urls
        product.finish         = data.get("finish", "")
        product.thickness      = data.get("thickness", "")
        product.dimensions     = data.get("dimensions", "")
        product.surface        = data.get("surface", "")
        product.product_type      = data.get("product_type", "")
        product.surface_category  = data.get("surface_category", "")
        product.application    = data.get("application", "")
        product.design_type    = data.get("design_type", "")
        product.color          = data.get("color", "")
        product.badge          = data.get("badge", "")
        product.accent_color   = data.get("accent_color", "") or "#85addc"
        product.meta_title     = data.get("meta_title", "")
        product.meta_description = data.get("meta_description", "")
        product.meta_keywords  = data.get("meta_keywords", "")
        image_ids = [i for i in data.getlist("image_ids") if i.strip()]
        has_images = bool(image_ids) or bool(image_urls)
        product.status         = _resolve_status(data, has_images, request)
        product.application_image_url = data.get("application_image_url", "")
        for f in _VISIBILITY_FIELDS:
            setattr(product, f, f in data)
        product.save()
        _save_product_images(product, data)
        # Related products — filter out blank strings
        rel_ids = [r for r in request.POST.getlist("related_products") if r.strip()]
        product.related_products.set(rel_ids)
        self.log_action("Updated product", product)
        messages.success(request, "Product updated.")
        return redirect("product_edit", pk=product.pk)


class ProductDeleteView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    def post(self, request, pk):
        product = get_object_or_404(Product, pk=pk)
        self.log_action("Deleted product", product)
        product.delete()
        messages.success(request, "Product deleted.")
        return redirect("product_list")


class ProductBulkDeleteView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    def post(self, request):
        ids = [i for i in request.POST.getlist("ids") if i.strip().isdigit()]
        qs = Product.objects.filter(pk__in=ids)
        count = qs.count()
        if count:
            names = ", ".join(qs.values_list("name", flat=True)[:5])
            qs.delete()
            self.log_action(f"Bulk deleted {count} product(s): {names}{'…' if count > 5 else ''}")
            messages.success(request, f"Deleted {count} product{'s' if count != 1 else ''}.")
        else:
            messages.error(request, "No products selected.")
        return redirect("product_list")


class ProductBulkStatusView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    def post(self, request):
        ids    = [i for i in request.POST.getlist("ids") if i.strip().isdigit()]
        status = request.POST.get("status")
        if not ids or status not in ("published", "draft"):
            messages.error(request, "Select at least one product and a valid status.")
            return redirect("product_list")
        qs    = Product.objects.filter(pk__in=ids)
        count = qs.count()
        qs.update(status=status)
        label = "published" if status == "published" else "set to draft"
        self.log_action(f"Bulk {label} {count} product(s)")
        messages.success(request, f"{count} product{'s' if count != 1 else ''} {label}.")
        return redirect("product_list")


class ProductPreviewView(ContentManagerRequiredMixin, View):
    def get(self, request, pk):
        product = get_object_or_404(Product, pk=pk)
        return render(request, "dashboard/preview/product.html", {
            "product":        product,
            "preview_title":  product.name,
            "preview_status": product.status,
            "edit_url":       f"/cms/products/{pk}/",
        })


# ── Bulk Export / Import (WooCommerce-style CSV) ──────────────────────────────

# Column order used for both export and the import template. `id` is optional on
# import (used to match an existing row); `sku` is the primary match key.
EXPORT_COLUMNS = [
    "id", "sku", "name", "slug", "status",
    "category", "collection",
    "short_description", "description",
    "features", "tech_specs",
    "finish", "thickness", "dimensions", "surface", "application",
    "product_type", "surface_category",
    "design_type", "color", "badge", "accent_color",
    "image_urls", "application_image_url", "texture_variants",
    "meta_title", "meta_description", "meta_keywords",
]

# Multi-value cells are separated by " | ".
_LIST_SEP = " | "


def _join_list(values):
    return _LIST_SEP.join(str(v).strip() for v in (values or []) if str(v).strip())


def _split_list(raw):
    if raw is None:
        return []
    return [p.strip() for p in str(raw).split("|") if p.strip()]


def _specs_to_cell(specs):
    """dict -> 'Key: value | Key: value'."""
    if not isinstance(specs, dict):
        return ""
    return _LIST_SEP.join(f"{k}: {v}" for k, v in specs.items())


def _cell_to_specs(raw):
    """'Key: value | Key: value' -> dict (order preserved)."""
    out = {}
    for part in _split_list(raw):
        if ":" in part:
            k, v = part.split(":", 1)
            k = k.strip()
            if k:
                out[k] = v.strip()
    return out


def _texture_to_cell(variants):
    """[{"label": ..., "image_url": ...}, ...] -> 'Label: url | Label: url'."""
    if not isinstance(variants, list):
        return ""
    return _LIST_SEP.join(
        f"{v.get('label', '')}: {v.get('image_url', '')}"
        for v in variants if isinstance(v, dict) and (v.get("label") or v.get("image_url"))
    )


def _cell_to_texture(raw):
    """'Label: url | Label: url' -> [{"label": ..., "image_url": ...}, ...]."""
    out = []
    for part in _split_list(raw):
        if ":" in part:
            label, url = part.split(":", 1)
            label, url = label.strip(), url.strip()
            if label or url:
                out.append({"label": label, "image_url": url})
    return out


def _product_to_row(p):
    return {
        "id":                p.pk,
        "sku":               p.sku,
        "name":              p.name,
        "slug":              p.slug,
        "status":            p.status,
        "category":          p.category.slug if p.category_id else "",
        "collection":        p.collection.slug if p.collection_id else "",
        "short_description": p.short_description,
        "description":       p.description,
        "features":          _join_list(p.features),
        "tech_specs":        _specs_to_cell(p.tech_specs),
        "finish":            p.finish,
        "thickness":         p.thickness,
        "dimensions":        p.dimensions,
        "surface":           p.surface,
        "product_type":      p.product_type,
        "surface_category":  p.surface_category,
        "application":       p.application,
        "design_type":       p.design_type,
        "color":             p.color,
        "badge":             p.badge,
        "accent_color":      p.accent_color,
        "image_urls":        _join_list(p.image_urls),
        "application_image_url": p.application_image_url,
        "texture_variants":  _texture_to_cell(p.texture_variants),
        "meta_title":        p.meta_title,
        "meta_description":  p.meta_description,
        "meta_keywords":     p.meta_keywords,
    }


def _parse_rows(file_obj, filename):
    """Parse an uploaded CSV or XLSX file into a list of dicts."""
    if filename.lower().endswith(".xlsx"):
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip() if c.value is not None else ""
                   for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            rows.append({h: ("" if v is None else str(v)) for h, v in zip(headers, row)})
        return rows
    text = file_obj.read().decode("utf-8-sig")
    return [dict(r) for r in csv.DictReader(io.StringIO(text))]


class ProductExportView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    """Download all (or the currently-filtered) products as a CSV."""

    def get(self, request):
        qs = Product.objects.select_related("category", "collection").order_by("sku")
        q      = request.GET.get("q", "")
        cat    = request.GET.get("category", "")
        status = request.GET.get("status", "")
        if q:
            qs = qs.filter(name__icontains=q) | qs.filter(sku__icontains=q)
        if cat:
            qs = qs.filter(category_id=cat)
        if status:
            qs = qs.filter(status=status)

        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="products-export.csv"'
        writer = csv.DictWriter(resp, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        count = 0
        for p in qs:
            writer.writerow(_product_to_row(p))
            count += 1
        self.log_action(f"Exported {count} product(s) to CSV")
        return resp


class ProductImportView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    template_name = "dashboard/products/bulk_import.html"

    def get(self, request):
        return render(request, self.template_name, {
            "active_nav": "products",
            "folders":     MediaFolder.objects.all(),
            "categories":  Category.objects.filter(status="published"),
            "collections": Collection.objects.filter(status="published"),
        })

    def post(self, request):
        action = request.POST.get("action", "preview")

        if action == "download_template":
            return self._download_template()

        if action == "folder_preview":
            return self._folder_preview(request)

        if action == "folder_commit":
            return self._folder_commit(request)

        create_missing = request.POST.get("create_missing_categories") == "on"

        if action == "commit":
            raw_rows = json.loads(request.POST.get("raw_rows", "[]"))
            create_missing = request.POST.get("create_missing_flag") == "1"
            created = updated = skipped = 0
            with transaction.atomic():
                for row in raw_rows:
                    result = self._apply_row(row, create_missing)
                    if result == "created":
                        created += 1
                    elif result == "updated":
                        updated += 1
                    else:
                        skipped += 1
            self.log_action(
                f"Bulk import: {created} created, {updated} updated, {skipped} skipped"
            )
            messages.success(
                request,
                f"Import complete — {created} created, {updated} updated, {skipped} skipped.",
            )
            return redirect("product_list")

        # action == "preview"
        uploaded = request.FILES.get("import_file")
        if not uploaded:
            messages.error(request, "No file uploaded.")
            return redirect("product_import")
        try:
            rows = _parse_rows(uploaded, uploaded.name)
        except Exception as exc:  # noqa: BLE001 — surface any parse failure to the user
            messages.error(request, f"Could not read file: {exc}")
            return redirect("product_import")

        preview_rows = []
        error_count = 0
        for i, row in enumerate(rows, start=2):
            info = self._validate_row(row, create_missing)
            info["row"] = i
            preview_rows.append(info)
            if info["errors"]:
                error_count += 1

        return render(request, self.template_name, {
            "preview_rows":   preview_rows,
            "error_count":    error_count,
            "ready_count":    len(preview_rows) - error_count,
            "raw_rows":       json.dumps(rows),
            "create_missing": create_missing,
            "active_nav":     "products",
        })

    # ── helpers ──────────────────────────────────────────────────────────────

    def _resolve_category(self, name, create_missing):
        key = (name or "").strip()
        if not key:
            return None, "Missing category"
        cat = (Category.objects.filter(slug__iexact=key).first()
               or Category.objects.filter(name__iexact=key).first())
        if cat:
            return cat, None
        if create_missing:
            return Category.objects.create(
                name=key, slug=slugify(key), status=Category.STATUS_PUBLISHED
            ), None
        return None, f"Unknown category '{key}'"

    def _resolve_collection(self, name):
        key = (name or "").strip()
        if not key:
            return None, None
        col = (Collection.objects.filter(slug__iexact=key).first()
               or Collection.objects.filter(name__iexact=key).first())
        if col:
            return col, None
        return None, f"Unknown collection '{key}'"

    def _validate_row(self, row, create_missing):
        errors = []
        sku = (row.get("sku") or "").strip()
        name = (row.get("name") or "").strip()
        raw_id = (str(row.get("id") or "")).strip()

        existing = None
        if raw_id.isdigit():
            existing = Product.objects.filter(pk=int(raw_id)).first()
            if not existing:
                errors.append(f"No product with id {raw_id}")
        if not existing and sku:
            existing = Product.objects.filter(sku__iexact=sku).first()

        if not sku:
            errors.append("Missing sku")
        if not existing and not name:
            errors.append("Missing name (required for new products)")

        cat_val = (row.get("category") or "").strip()
        if cat_val or not existing:
            _, cat_err = self._resolve_category(cat_val, create_missing)
            # When updating an existing row a blank category keeps the current one.
            if cat_err and (cat_val or not existing):
                errors.append(cat_err)

        _, col_err = self._resolve_collection(row.get("collection"))
        if col_err:
            errors.append(col_err)

        for field in ("design_type", "color", "badge"):
            val = (row.get(field) or "").strip()
            if val and val.lower() not in _attr_lookup(field):
                errors.append(f"Invalid {field} '{val}'")

        status = (row.get("status") or "").strip().lower()
        if status and status not in ("draft", "published"):
            errors.append(f"Invalid status '{status}'")

        return {
            "sku":    sku,
            "name":   name or (existing.name if existing else ""),
            "action": "skip" if errors else ("update" if existing else "create"),
            "errors": errors,
        }

    def _apply_row(self, row, create_missing):
        info = self._validate_row(row, create_missing)
        if info["errors"]:
            return "skipped"

        sku = info["sku"]
        raw_id = (str(row.get("id") or "")).strip()
        product = None
        if raw_id.isdigit():
            product = Product.objects.filter(pk=int(raw_id)).first()
        if not product:
            product = Product.objects.filter(sku__iexact=sku).first()
        is_new = product is None
        if is_new:
            product = Product(sku=sku)

        def has(col):
            return col in row and str(row.get(col)).strip() != ""

        if has("name"):
            product.name = row["name"].strip()
        if has("sku"):
            product.sku = sku
        if has("slug"):
            product.slug = slugify(row["slug"].strip())
        elif is_new:
            product.slug = ""  # model.save() will generate

        cat_val = (row.get("category") or "").strip()
        if cat_val or is_new:
            cat, _ = self._resolve_category(cat_val, create_missing)
            if cat:
                product.category = cat
        if "collection" in row:
            col_val = (row.get("collection") or "").strip()
            if col_val:
                col, _ = self._resolve_collection(col_val)
                product.collection = col
            else:
                product.collection = None

        text_fields = [
            "short_description", "description", "finish", "thickness",
            "dimensions", "surface", "product_type", "surface_category",
            "application", "accent_color", "application_image_url",
            "meta_title", "meta_description", "meta_keywords",
        ]
        for f in text_fields:
            if f in row:
                setattr(product, f, (row.get(f) or "").strip())

        if has("design_type"):
            product.design_type = _attr_lookup("design_type")[row["design_type"].strip().lower()]
        if has("color"):
            product.color = _attr_lookup("color")[row["color"].strip().lower()]
        if has("badge"):
            product.badge = _attr_lookup("badge")[row["badge"].strip().lower()]
        if not product.accent_color:
            product.accent_color = "#85addc"

        if "features" in row:
            product.features = _split_list(row.get("features"))
        if "tech_specs" in row:
            product.tech_specs = _cell_to_specs(row.get("tech_specs"))
        if "image_urls" in row:
            product.image_urls = _split_list(row.get("image_urls"))
        if "texture_variants" in row:
            product.texture_variants = _cell_to_texture(row.get("texture_variants"))

        if has("status"):
            product.status = row["status"].strip().lower()
        elif is_new and not product.status:
            product.status = "draft"

        if product.status == Product.STATUS_PUBLISHED:
            has_images = bool(product.image_urls) or (
                not is_new and product.product_images.filter(role=ProductImage.ROLE_GALLERY).exists()
            )
            if not has_images:
                product.status = Product.STATUS_DRAFT

        try:
            with transaction.atomic():
                product.save()
        except IntegrityError:
            return "skipped"
        return "created" if is_new else "updated"

    def _download_template(self):
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="products-import-template.csv"'
        writer = csv.DictWriter(resp, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        writer.writerow({
            "id": "", "sku": "LAM-001", "name": "Sample Laminate",
            "slug": "", "status": "draft",
            "category": "laminates", "collection": "",
            "short_description": "One-line summary",
            "description": "<p>Full rich-text description</p>",
            "features": "Scratch resistant | Heat resistant | Anti-fingerprint",
            "tech_specs": "Thickness: 1.0mm | Warranty: 10 years",
            "finish": "High Gloss", "thickness": "1.0mm",
            "dimensions": "8ft x 4ft (2440 x 1220mm)",
            "surface": "Decorative Laminate",
            "application": "Cabinets, Wardrobes, Wall Panels",
            "design_type": "Wood", "color": "Brown", "badge": "New",
            "accent_color": "#85addc",
            "image_urls": "https://example.com/a.jpg | https://example.com/b.jpg",
            "application_image_url": "https://example.com/applied-in-room.jpg",
            "texture_variants": "Fluted: https://example.com/fluted.jpg | Glossy: https://example.com/glossy.jpg",
            "meta_title": "Sample Laminate | Sanish",
            "meta_description": "Buy sample laminate…",
            "meta_keywords": "laminate, wood laminate",
        })
        return resp

    # ── Auto-attach images from a Media Library folder ─────────────────────
    #
    # Same convention as the `import_sheet_range` management command: every
    # file is named by its product code (SKU), with an optional suffix marking
    # its role on the product page:
    #   <code>-preview / -applied / -pw       -> the "applied in a room" shot
    #   <code>-texture-<name>                 -> a texture-variant thumbnail,
    #                                             e.g. "6005-5-texture-fluted.jpg"
    #                                             -> label "Fluted"
    #   <code> (no suffix)                    -> a plain gallery/swatch image
    # The code becomes the product name when a new product is created.

    _FOLDER_PREVIEW_SUFFIX_RE = re.compile(r"[-_ ](preview|applied|pw)$", re.I)
    _FOLDER_TEXTURE_SUFFIX_RE = re.compile(r"[-_ ]texture[-_ ]([a-z0-9]+)$", re.I)
    _FOLDER_ROLE_ORDER = {
        ProductImage.ROLE_GALLERY: 0, ProductImage.ROLE_APPLICATION: 1, ProductImage.ROLE_TEXTURE: 2,
    }

    def _classify_folder_asset(self, stem):
        """Filename stem -> (product code, ProductImage role, texture label)."""
        m = self._FOLDER_TEXTURE_SUFFIX_RE.search(stem)
        if m:
            label = m.group(1).replace("-", " ").replace("_", " ").title()
            code = self._FOLDER_TEXTURE_SUFFIX_RE.sub("", stem).strip()
            return code, ProductImage.ROLE_TEXTURE, label
        if self._FOLDER_PREVIEW_SUFFIX_RE.search(stem):
            code = self._FOLDER_PREVIEW_SUFFIX_RE.sub("", stem).strip()
            return code, ProductImage.ROLE_APPLICATION, ""
        return stem, ProductImage.ROLE_GALLERY, ""

    def _group_folder_assets(self, folder):
        """MediaFolder -> {code: [(asset, role, label), ...]}, gallery image(s) first."""
        groups = {}
        for asset in sorted(folder.assets.all(), key=lambda a: a.filename):
            stem, _ext = os.path.splitext(asset.filename)
            code, role, label = self._classify_folder_asset(stem)
            if not code:
                continue
            groups.setdefault(code, []).append((asset, role, label))
        for items in groups.values():
            items.sort(key=lambda t: (self._FOLDER_ROLE_ORDER[t[1]], t[0].filename))
        return groups

    def _folder_preview(self, request):
        folder = get_object_or_404(MediaFolder, pk=request.POST.get("folder_id"))
        category_id = request.POST.get("folder_category") or ""
        collection_id = request.POST.get("folder_collection") or ""
        create_missing = request.POST.get("folder_create_missing") == "on"
        groups = self._group_folder_assets(folder)

        rows = []
        for code in sorted(groups):
            assets = groups[code]
            product = (Product.objects.filter(sku__iexact=code).first()
                       or Product.objects.filter(slug=slugify(code)).first())
            if product:
                action = "update"
            elif create_missing and category_id:
                action = "create"
            else:
                action = "skip"
            rows.append({
                "code": code,
                "product": product,
                "action": action,
                "images": [a for a, _role, _label in assets],
            })

        return render(request, self.template_name, {
            "active_nav":   "products",
            "folder_rows":  rows,
            "folder":       folder,
            "folder_category_id":   category_id,
            "folder_collection_id": collection_id,
            "folder_create_missing": create_missing,
            "folders":      MediaFolder.objects.all(),
            "categories":   Category.objects.filter(status="published"),
            "collections":  Collection.objects.filter(status="published"),
        })

    def _folder_commit(self, request):
        folder = get_object_or_404(MediaFolder, pk=request.POST.get("folder_id"))
        category_id = request.POST.get("folder_category") or None
        collection_id = request.POST.get("folder_collection") or None
        category = Category.objects.filter(pk=category_id).first() if category_id else None
        collection = Collection.objects.filter(pk=collection_id).first() if collection_id else None
        create_missing = request.POST.get("folder_create_missing_flag") == "1"
        codes = request.POST.getlist("codes")
        groups = self._group_folder_assets(folder)

        created = updated = skipped = attached = 0
        with transaction.atomic():
            for code in codes:
                assets = groups.get(code)
                if not assets:
                    skipped += 1
                    continue

                product = (Product.objects.filter(sku__iexact=code).first()
                           or Product.objects.filter(slug=slugify(code)).first())
                is_new = product is None
                if is_new:
                    if not create_missing or not category:
                        skipped += 1
                        continue
                    product = Product(
                        name=code, sku=code, slug=slugify(code),
                        category=category, collection=collection,
                        status="draft", accent_color="#85addc",
                    )
                    product.save()
                    created += 1
                else:
                    updated += 1

                existing_asset_ids = set(
                    product.product_images.values_list("asset_id", flat=True)
                )
                next_pos = product.product_images.count()
                for asset, role, label in assets:
                    if asset.id in existing_asset_ids:
                        continue
                    ProductImage.objects.create(
                        product=product, asset=asset, position=next_pos, role=role, label=label,
                    )
                    next_pos += 1
                    attached += 1

        self.log_action(
            f"Folder image import ('{folder.name}'): {created} created, {updated} matched, "
            f"{attached} image(s) attached, {skipped} skipped"
        )
        messages.success(
            request,
            f"Folder import complete — {created} created, {updated} matched, "
            f"{attached} image(s) attached, {skipped} skipped.",
        )


class AttributeOptionAddView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    """Add a new value to a managed dropdown (Design Type, Colour, Badge, Finish)."""

    def post(self, request, attribute):
        if attribute not in MANAGED_ATTRIBUTES:
            return JsonResponse({"ok": False, "error": "Unknown attribute"}, status=404)
        try:
            payload = json.loads(request.body or "{}")
        except json.JSONDecodeError:
            return JsonResponse({"ok": False, "error": "Bad payload"}, status=400)
        value = (payload.get("value") or "").strip()
        if not value:
            return JsonResponse({"ok": False, "error": "Value is required"}, status=400)

        existing = ProductAttributeOption.objects.filter(
            attribute=attribute, value__iexact=value
        ).first()
        if existing:
            return JsonResponse({"ok": True, "id": existing.id, "value": existing.value})

        next_pos = (
            ProductAttributeOption.objects.filter(attribute=attribute)
            .count()
        )
        option = ProductAttributeOption.objects.create(
            attribute=attribute, value=value, position=next_pos
        )
        self.log_action(f"Added {attribute} option “{value}”", option)
        return JsonResponse({"ok": True, "id": option.id, "value": option.value})


class AttributeOptionDeleteView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    """Remove a value from a managed dropdown."""

    def post(self, request, attribute):
        if attribute not in MANAGED_ATTRIBUTES:
            return JsonResponse({"ok": False, "error": "Unknown attribute"}, status=404)
        try:
            payload = json.loads(request.body or "{}")
        except json.JSONDecodeError:
            return JsonResponse({"ok": False, "error": "Bad payload"}, status=400)
        option = ProductAttributeOption.objects.filter(
            attribute=attribute, pk=payload.get("id")
        ).first()
        if not option:
            return JsonResponse({"ok": False, "error": "Not found"}, status=404)
        value = option.value
        option.delete()
        self.log_action(f"Removed {attribute} option “{value}”", None)
        return JsonResponse({"ok": True})


class AttributeOptionReorderView(ContentManagerRequiredMixin, LoggedActionMixin, View):
    """Persist the drag-to-reorder sequence for a managed dropdown."""

    def post(self, request, attribute):
        if attribute not in MANAGED_ATTRIBUTES:
            return JsonResponse({"ok": False, "error": "Unknown attribute"}, status=404)
        try:
            order = json.loads(request.body or "{}").get("order", [])
        except json.JSONDecodeError:
            return JsonResponse({"ok": False, "error": "Bad payload"}, status=400)
        for pos, pk in enumerate(order):
            ProductAttributeOption.objects.filter(pk=pk, attribute=attribute).update(position=pos)
        return JsonResponse({"ok": True})
        return redirect("product_list")
